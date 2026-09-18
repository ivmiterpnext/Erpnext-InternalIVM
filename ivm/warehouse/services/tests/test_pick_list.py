"""Integration tests for ivm.warehouse.services.pick_list"""

import frappe
from erpnext.tests.utils import ERPNextTestSuite
from erpnext.stock.doctype.item.test_item import make_item
from erpnext.stock.doctype.stock_entry.stock_entry_utils import make_stock_entry
from xlsxwriter.url import Url

from ivm.warehouse.services.pick_list import (
    _autofit_column_width,
    _build_item_link,
    _build_pick_list_excel_payload,
    _get_related_project,
    _pixels_to_character_width,
    _resolve_pick_list_label,
    add_item_to_pick_list,
    clear_pick_list_items,
    create_pick_list,
    delete_draft_pick_list,
    export_pick_list_cost_excel,
    get_pick_list_cost_rows,
    remove_pick_list_item,
    serialize_pick_list,
    submit_pick_list,
    update_pick_list_item_qty,
)

COMPANY = "_Test Company"
WAREHOUSE = "_Test Warehouse - _TC"


def _seed_stock(item_code, warehouse=WAREHOUSE, qty=100, rate=10):
    """Create stock via Material Receipt so Bin.actual_qty is populated."""
    se = make_stock_entry(
        item_code=item_code,
        qty=qty,
        basic_rate=rate,
        to_warehouse=warehouse,
        company=COMPANY,
        stock_entry_type="Material Receipt",
    )
    se.submit()


class TestCreatePickList(ERPNextTestSuite):
    """create_pick_list / delete_draft_pick_list"""
    
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        # Ensure test company and warehouse exist
        if not frappe.db.exists("Company", "_Test Company"):
            frappe.get_doc({
                "doctype": "Company",
                "company_name": "_Test Company",
                "abbr": "_TC",
                "country": "United States",
                "default_currency": "USD"
            }).insert(ignore_permissions=True)
        
        if not frappe.db.exists("Warehouse", "_Test Warehouse - _TC"):
            frappe.get_doc({
                "doctype": "Warehouse",
                "warehouse_name": "_Test Warehouse",
                "company": "_Test Company"
            }).insert(ignore_permissions=True)

    def test_create_returns_draft_name(self):
        name = create_pick_list(COMPANY)
        self.assertTrue(name)
        pl = frappe.get_doc("Pick List", name)
        self.assertEqual(pl.docstatus, 0)
        self.assertEqual(pl.company, COMPANY)
        self.assertEqual(pl.purpose, "Material Transfer")

    def test_delete_draft(self):
        name = create_pick_list(COMPANY)
        delete_draft_pick_list(name)
        self.assertFalse(frappe.db.exists("Pick List", name))

    def test_delete_submitted_raises(self):
        item = make_item()
        _seed_stock(item.name)
        name = create_pick_list(COMPANY)
        add_item_to_pick_list(name, item.name, WAREHOUSE, 1)
        pl = frappe.get_doc("Pick List", name)
        pl.submit()
        with self.assertRaises(frappe.ValidationError):
            delete_draft_pick_list(name)


class TestAddItem(ERPNextTestSuite):
    """add_item_to_pick_list"""

    def test_add_new_item(self):
        item = make_item()
        _seed_stock(item.name)
        name = create_pick_list(COMPANY)
        result = add_item_to_pick_list(name, item.name, WAREHOUSE, 5)
        self.assertIn("row_name", result)
        self.assertEqual(result["qty"], 5)

    def test_add_increments_existing(self):
        item = make_item()
        _seed_stock(item.name)
        name = create_pick_list(COMPANY)
        r1 = add_item_to_pick_list(name, item.name, WAREHOUSE, 3)
        r2 = add_item_to_pick_list(name, item.name, WAREHOUSE, 2)
        self.assertEqual(r2["qty"], 5)
        self.assertEqual(r1["row_name"], r2["row_name"])

    def test_string_qty_coerced_to_float(self):
        item = make_item()
        _seed_stock(item.name)
        name = create_pick_list(COMPANY)
        result = add_item_to_pick_list(name, item.name, WAREHOUSE, "7")
        self.assertEqual(result["qty"], 7)

    def test_fetches_item_name_and_uom_when_omitted(self):
        item = make_item()
        _seed_stock(item.name)
        name = create_pick_list(COMPANY)
        add_item_to_pick_list(name, item.name, WAREHOUSE, 1)
        pl = frappe.get_doc("Pick List", name)
        loc = pl.locations[0]
        self.assertTrue(loc.item_name)
        self.assertTrue(loc.uom)

    def test_add_to_submitted_pick_list_raises(self):
        item = make_item()
        _seed_stock(item.name)
        name = create_pick_list(COMPANY)
        add_item_to_pick_list(name, item.name, WAREHOUSE, 1)
        frappe.get_doc("Pick List", name).submit()
        with self.assertRaises(frappe.ValidationError):
            add_item_to_pick_list(name, item.name, WAREHOUSE, 1)


class TestRemoveItem(ERPNextTestSuite):
    """remove_pick_list_item — including Priority 1 DoesNotExistError fix"""

    def test_remove_existing_row(self):
        item = make_item()
        _seed_stock(item.name)
        name = create_pick_list(COMPANY)
        r = add_item_to_pick_list(name, item.name, WAREHOUSE, 5)
        result = remove_pick_list_item(name, r["row_name"])
        self.assertTrue(result["success"])
        pl = frappe.get_doc("Pick List", name)
        self.assertEqual(len(pl.locations), 0)

    def test_remove_nonexistent_row_raises_does_not_exist(self):
        name = create_pick_list(COMPANY)
        with self.assertRaises(frappe.DoesNotExistError):
            remove_pick_list_item(name, "nonexistent-row-id")


class TestUpdateItemQty(ERPNextTestSuite):
    """update_pick_list_item_qty — including Priority 1 DoesNotExistError fix"""

    def test_update_qty(self):
        item = make_item()
        _seed_stock(item.name)
        name = create_pick_list(COMPANY)
        r = add_item_to_pick_list(name, item.name, WAREHOUSE, 5)
        update_pick_list_item_qty(name, r["row_name"], 10)
        pl = frappe.get_doc("Pick List", name)
        loc = pl.locations[0]
        self.assertEqual(loc.qty, 10)
        self.assertEqual(loc.picked_qty, 10)

    def test_update_with_string_qty(self):
        item = make_item()
        _seed_stock(item.name)
        name = create_pick_list(COMPANY)
        r = add_item_to_pick_list(name, item.name, WAREHOUSE, 5)
        update_pick_list_item_qty(name, r["row_name"], "3")
        pl = frappe.get_doc("Pick List", name)
        self.assertEqual(pl.locations[0].qty, 3)

    def test_update_nonexistent_row_raises_does_not_exist(self):
        name = create_pick_list(COMPANY)
        with self.assertRaises(frappe.DoesNotExistError):
            update_pick_list_item_qty(name, "nonexistent-row-id", 5)


class TestClearItems(ERPNextTestSuite):
    """clear_pick_list_items"""

    def test_clear_removes_all_rows(self):
        item_a = make_item()
        item_b = make_item()
        _seed_stock(item_a.name)
        _seed_stock(item_b.name)
        name = create_pick_list(COMPANY)
        add_item_to_pick_list(name, item_a.name, WAREHOUSE, 3)
        add_item_to_pick_list(name, item_b.name, WAREHOUSE, 2)
        clear_pick_list_items(name)
        pl = frappe.get_doc("Pick List", name)
        self.assertEqual(len(pl.locations), 0)

    def test_clear_empty_pick_list_succeeds(self):
        name = create_pick_list(COMPANY)
        result = clear_pick_list_items(name)
        self.assertTrue(result["success"])


class TestSerializePickList(ERPNextTestSuite):
    """serialize_pick_list"""

    def test_draft_serialization_shape(self):
        item = make_item()
        _seed_stock(item.name, qty=50)
        name = create_pick_list(COMPANY)
        add_item_to_pick_list(name, item.name, WAREHOUSE, 5)
        pl = frappe.get_doc("Pick List", name)
        data = serialize_pick_list(pl)
        self.assertEqual(data["pick_list"], name)
        self.assertFalse(data["submitted"])
        self.assertIsNone(data["stock_entry"])
        self.assertEqual(len(data["items"]), 1)
        row = data["items"][0]
        self.assertEqual(row["item_code"], item.name)
        self.assertEqual(row["qty"], 5)
        self.assertEqual(row["available_qty"], 50)
        self.assertIn("row_name", row)
        self.assertIn("uom", row)

    def test_draft_available_qty_is_live(self):
        """Draft serialization should query current Bin qty, not stale stock_qty."""
        item = make_item()
        _seed_stock(item.name, qty=20)
        name = create_pick_list(COMPANY)
        add_item_to_pick_list(name, item.name, WAREHOUSE, 5)
        # Add more stock after the pick list row was created
        _seed_stock(item.name, qty=30)
        pl = frappe.get_doc("Pick List", name)
        data = serialize_pick_list(pl)
        self.assertEqual(data["items"][0]["available_qty"], 50)


class TestSubmitPickList(ERPNextTestSuite):
    """submit_pick_list"""

    def test_submit_creates_stock_entry(self):
        item = make_item()
        _seed_stock(item.name)
        name = create_pick_list(COMPANY)
        add_item_to_pick_list(name, item.name, WAREHOUSE, 5)
        result = submit_pick_list(name, target_warehouse=WAREHOUSE)
        self.assertIn("stock_entry", result)
        self.assertTrue(result["stock_entry"])
        se = frappe.get_doc("Stock Entry", result["stock_entry"])
        self.assertEqual(se.docstatus, 0)  # draft Stock Entry

    def test_submit_with_target_warehouse(self):
        item = make_item()
        _seed_stock(item.name)
        name = create_pick_list(COMPANY)
        add_item_to_pick_list(name, item.name, WAREHOUSE, 5)
        target = "_Test Warehouse 1 - _TC"
        result = submit_pick_list(name, target_warehouse=target)
        se = frappe.get_doc("Stock Entry", result["stock_entry"])
        for row in se.items:
            self.assertEqual(row.t_warehouse, target)

    def test_submit_without_target_warehouse_raises(self):
        """A Pick List with no parent_warehouse and no target_warehouse arg
        should fail at the Stock Entry validation layer, since there is no
        destination warehouse for the transfer."""
        item = make_item()
        _seed_stock(item.name)
        name = create_pick_list(COMPANY)
        add_item_to_pick_list(name, item.name, WAREHOUSE, 5)
        with self.assertRaises(frappe.ValidationError):
            submit_pick_list(name)

    def test_submit_links_warehouse_request(self):
        item = make_item()
        _seed_stock(item.name)
        name = create_pick_list(COMPANY)
        add_item_to_pick_list(name, item.name, WAREHOUSE, 5)
        # Create a Warehouse Request pointing at this Pick List
        wr = frappe.get_doc({
            "doctype": "Warehouse Request",
            "pick_list": name,
            "request_reason": "Build Machine",
            "subject": "Test WR",
        })
        wr.insert(ignore_permissions=True)
        result = submit_pick_list(name, target_warehouse=WAREHOUSE)
        se = frappe.get_doc("Stock Entry", result["stock_entry"])
        self.assertEqual(se.custom_warehouse_request, wr.name)

    def test_submit_already_submitted_raises(self):
        item = make_item()
        _seed_stock(item.name)
        name = create_pick_list(COMPANY)
        add_item_to_pick_list(name, item.name, WAREHOUSE, 5)
        frappe.get_doc("Pick List", name).submit()
        with self.assertRaises(frappe.ValidationError):
            submit_pick_list(name)


class TestGetPickListCostRows(ERPNextTestSuite):
    """get_pick_list_cost_rows"""

    def test_uses_bin_valuation_rate(self):
        item = make_item()
        _seed_stock(item.name, qty=10, rate=25)
        name = create_pick_list(COMPANY)
        add_item_to_pick_list(name, item.name, WAREHOUSE, 4)
        rows = get_pick_list_cost_rows(name)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["item_code"], item.name)
        self.assertEqual(rows[0]["qty"], 4)
        self.assertEqual(rows[0]["rate"], 25)
        self.assertEqual(rows[0]["amount"], 100)

    def test_falls_back_to_item_valuation_rate_when_no_bin(self):
        item = make_item()
        # No stock seeded for this item at WAREHOUSE — no Bin record exists.
        frappe.db.set_value("Item", item.name, "valuation_rate", 42)
        name = create_pick_list(COMPANY)
        pl = frappe.get_doc("Pick List", name)
        pl.append("locations", {
            "item_code": item.name,
            "item_name": item.item_name,
            "warehouse": WAREHOUSE,
            "qty": 2,
            "picked_qty": 2,
            "uom": item.stock_uom,
            "stock_uom": item.stock_uom,
            "conversion_factor": 1,
        })
        pl.flags.ignore_validate = True
        pl.save()
        rows = get_pick_list_cost_rows(name)
        self.assertEqual(rows[0]["rate"], 42)
        self.assertEqual(rows[0]["amount"], 84)

    def test_multiple_items_row_shape_and_fields(self):
        item_a = make_item()
        item_b = make_item()
        _seed_stock(item_a.name, qty=5, rate=10)
        _seed_stock(item_b.name, qty=5, rate=20)
        name = create_pick_list(COMPANY)
        add_item_to_pick_list(name, item_a.name, WAREHOUSE, 2)
        add_item_to_pick_list(name, item_b.name, WAREHOUSE, 3)
        rows = get_pick_list_cost_rows(name)
        self.assertEqual(len(rows), 2)
        total = sum(r["amount"] for r in rows)
        self.assertEqual(total, 2 * 10 + 3 * 20)
        for row in rows:
            self.assertIn("item_code", row)
            self.assertIn("item_name", row)
            self.assertIn("warehouse", row)
            self.assertIn("uom", row)

    def test_empty_pick_list_returns_empty_rows(self):
        name = create_pick_list(COMPANY)
        rows = get_pick_list_cost_rows(name)
        self.assertEqual(rows, [])


class TestGetRelatedProject(ERPNextTestSuite):
    """_get_related_project"""

    def test_returns_project_when_warehouse_request_links_pick_list(self):
        name = create_pick_list(COMPANY)
        project = frappe.get_doc({
            "doctype": "Project",
            "project_name": "Test Project For Pick List Cost Export",
            "company": COMPANY,
        }).insert(ignore_permissions=True)
        wr = frappe.get_doc({
            "doctype": "Warehouse Request",
            "pick_list": name,
            "related_project": project.name,
            "request_reason": "Build Machine",
            "subject": "Test WR",
        })
        wr.insert(ignore_permissions=True)
        project_id, project_name = _get_related_project(name)
        self.assertEqual(project_id, project.name)
        self.assertEqual(project_name, project.project_name)

    def test_returns_none_when_no_warehouse_request_references_pick_list(self):
        name = create_pick_list(COMPANY)
        project_id, project_name = _get_related_project(name)
        self.assertIsNone(project_id)
        self.assertIsNone(project_name)

    def test_returns_none_when_warehouse_request_has_no_related_project(self):
        name = create_pick_list(COMPANY)
        wr = frappe.get_doc({
            "doctype": "Warehouse Request",
            "pick_list": name,
            "request_reason": "Build Machine",
            "subject": "Test WR No Project",
        })
        wr.insert(ignore_permissions=True)
        project_id, project_name = _get_related_project(name)
        self.assertIsNone(project_id)
        self.assertIsNone(project_name)


class TestResolvePickListLabel(ERPNextTestSuite):
    """_resolve_pick_list_label"""

    def test_prefers_machine_names(self):
        name = create_pick_list(COMPANY)
        wr = frappe.get_doc({
            "doctype": "Warehouse Request",
            "pick_list": name,
            "request_reason": "Build Machine",
            "subject": "Test WR",
            "machine_names": "ED01748L",
            "machine_name": "ED01748L-primary-ignored",
            "related_project": "",
        })
        wr.insert(ignore_permissions=True)
        self.assertEqual(_resolve_pick_list_label(name), "ED01748L")

    def test_falls_back_to_machine_name_when_machine_names_blank(self):
        name = create_pick_list(COMPANY)
        wr = frappe.get_doc({
            "doctype": "Warehouse Request",
            "pick_list": name,
            "request_reason": "Build Machine",
            "subject": "Test WR",
            "machine_name": "ED01644V",
        })
        wr.insert(ignore_permissions=True)
        self.assertEqual(_resolve_pick_list_label(name), "ED01644V")

    def test_falls_back_to_related_project_when_no_machine_name(self):
        name = create_pick_list(COMPANY)
        project = frappe.get_doc({
            "doctype": "Project",
            "project_name": "Test Project For Pick List Label",
            "company": COMPANY,
        }).insert(ignore_permissions=True)
        wr = frappe.get_doc({
            "doctype": "Warehouse Request",
            "pick_list": name,
            "request_reason": "Shipping Request",
            "subject": "Ship Parts",
            "related_project": project.name,
        })
        wr.insert(ignore_permissions=True)
        self.assertEqual(_resolve_pick_list_label(name), project.name)

    def test_falls_back_to_pick_list_name_when_no_warehouse_request(self):
        name = create_pick_list(COMPANY)
        self.assertEqual(_resolve_pick_list_label(name), name)

    def test_falls_back_to_pick_list_name_when_warehouse_request_has_no_label_fields(self):
        name = create_pick_list(COMPANY)
        wr = frappe.get_doc({
            "doctype": "Warehouse Request",
            "pick_list": name,
            "request_reason": "Shipping Request",
            "subject": "Ship Parts",
        })
        wr.insert(ignore_permissions=True)
        self.assertEqual(_resolve_pick_list_label(name), name)

    def test_sanitizes_illegal_filename_characters(self):
        name = create_pick_list(COMPANY)
        wr = frappe.get_doc({
            "doctype": "Warehouse Request",
            "pick_list": name,
            "request_reason": "Build Machine",
            "subject": "Test WR",
            "machine_names": 'ED0174/8L:*?"<>|',
        })
        wr.insert(ignore_permissions=True)
        self.assertEqual(_resolve_pick_list_label(name), "ED01748L")


class TestBuildPickListExcelPayload(ERPNextTestSuite):
    """_build_pick_list_excel_payload"""

    def test_header_row_matches_template(self):
        name = create_pick_list(COMPANY)
        payload = _build_pick_list_excel_payload(name)
        self.assertEqual(
            payload["data"][0],
            ["Inventory Part ", None, "Qty Picked", None, "Price", None, "Total"],
        )

    def test_item_rows_and_total_formula(self):
        item_a = make_item()
        item_b = make_item()
        _seed_stock(item_a.name, qty=5, rate=10)
        _seed_stock(item_b.name, qty=5, rate=20)
        name = create_pick_list(COMPANY)
        add_item_to_pick_list(name, item_a.name, WAREHOUSE, 2)
        add_item_to_pick_list(name, item_b.name, WAREHOUSE, 3)
        payload = _build_pick_list_excel_payload(name)
        data = payload["data"]

        self.assertIsInstance(data[1][0], Url)
        self.assertEqual(data[1][0].text, f"{item_a.name}: {item_a.item_name}")
        self.assertEqual(data[1][2], 2)
        self.assertEqual(data[1][4], 10)
        self.assertEqual(data[1][6], "=C2*E2")

        self.assertIsInstance(data[2][0], Url)
        self.assertEqual(data[2][0].text, f"{item_b.name}: {item_b.item_name}")
        self.assertEqual(data[2][2], 3)
        self.assertEqual(data[2][4], 20)
        self.assertEqual(data[2][6], "=C3*E3")

        # index 3 = blank spacer row, index 4 = total row
        self.assertEqual(data[3], [None] * 7)
        self.assertEqual(data[4], [None, None, None, None, None, None, "=SUM(G2:G3)"])

    def test_empty_pick_list_has_no_sum_formula(self):
        name = create_pick_list(COMPANY)
        payload = _build_pick_list_excel_payload(name)
        data = payload["data"]
        self.assertEqual(len(data), 3)  # header + blank spacer + blank total row
        self.assertEqual(data[1], [None] * 7)
        self.assertEqual(data[2], [None] * 7)

    def test_price_and_total_column_widths_are_fixed(self):
        name = create_pick_list(COMPANY)
        payload = _build_pick_list_excel_payload(name)
        widths = payload["column_widths"]
        self.assertEqual(widths[4], 11)
        self.assertEqual(widths[6], 12)
        self.assertEqual(widths[2], 10.14)
        self.assertIsNone(widths[1])
        self.assertIsNone(widths[3])
        self.assertIsNone(widths[5])
        self.assertIn("column_styles", payload["styles"])
        self.assertEqual(payload["styles"]["column_styles"][4], [1])
        self.assertEqual(payload["styles"]["column_styles"][6], [1])

    def test_inventory_part_width_floor_when_descriptions_are_short(self):
        item = make_item()
        _seed_stock(item.name, qty=5, rate=10)
        name = create_pick_list(COMPANY)
        add_item_to_pick_list(name, item.name, WAREHOUSE, 2)
        payload = _build_pick_list_excel_payload(name)
        # short item name/code shouldn't push the width below the template floor
        self.assertEqual(payload["column_widths"][0], 57.86)

    def test_inventory_part_width_grows_for_long_descriptions(self):
        long_name = "A" * 140
        item = make_item(properties={"item_name": long_name})
        _seed_stock(item.name, qty=5, rate=10)
        name = create_pick_list(COMPANY)
        add_item_to_pick_list(name, item.name, WAREHOUSE, 2, item_name=long_name)
        payload = _build_pick_list_excel_payload(name)
        self.assertGreater(payload["column_widths"][0], 57.86)

    def test_inventory_part_width_floor_when_pick_list_empty(self):
        name = create_pick_list(COMPANY)
        payload = _build_pick_list_excel_payload(name)
        self.assertEqual(payload["column_widths"][0], 57.86)

    def test_filename_uses_resolved_label(self):
        name = create_pick_list(COMPANY)
        wr = frappe.get_doc({
            "doctype": "Warehouse Request",
            "pick_list": name,
            "request_reason": "Build Machine",
            "subject": "Test WR",
            "machine_names": "ED01748L",
        })
        wr.insert(ignore_permissions=True)
        payload = _build_pick_list_excel_payload(name)
        self.assertEqual(payload["filename"], "ED01748L - Pick List")

    def test_last_item_row_has_bottom_border_across_all_columns(self):
        item_a = make_item()
        item_b = make_item()
        _seed_stock(item_a.name, qty=5, rate=10)
        _seed_stock(item_b.name, qty=5, rate=20)
        name = create_pick_list(COMPANY)
        add_item_to_pick_list(name, item_a.name, WAREHOUSE, 2)
        add_item_to_pick_list(name, item_b.name, WAREHOUSE, 3)
        payload = _build_pick_list_excel_payload(name)

        border_style_id = payload["styles"]["styles"].index({"bottom": 1})
        hyperlink_style_id = payload["styles"]["styles"].index(
            {"font_color": "blue", "underline": 1, "align": "left", "valign": "vcenter"}
        )
        last_item_row_idx = 2  # header=0, item_a=1, item_b=2

        cell_styles = payload["styles"]["cell_styles"]
        self.assertEqual(cell_styles[(last_item_row_idx, 0)], [hyperlink_style_id, border_style_id])
        for col_idx in range(1, 7):
            self.assertEqual(cell_styles[(last_item_row_idx, col_idx)], [border_style_id])

        # first item row: hyperlink style present, but no border
        self.assertEqual(cell_styles[(1, 0)], [hyperlink_style_id])
        self.assertNotIn((1, 1), cell_styles)

    def test_empty_pick_list_has_no_border_cell_styles(self):
        name = create_pick_list(COMPANY)
        payload = _build_pick_list_excel_payload(name)
        self.assertEqual(payload["styles"]["cell_styles"], {})

    def test_last_item_row_idx_present_when_rows_exist(self):
        item = make_item()
        _seed_stock(item.name, qty=5, rate=10)
        name = create_pick_list(COMPANY)
        add_item_to_pick_list(name, item.name, WAREHOUSE, 2)
        payload = _build_pick_list_excel_payload(name)
        self.assertEqual(payload["last_item_row_idx"], 1)

    def test_last_item_row_idx_none_when_empty(self):
        name = create_pick_list(COMPANY)
        payload = _build_pick_list_excel_payload(name)
        self.assertIsNone(payload["last_item_row_idx"])


class TestExportPickListCostExcel(ERPNextTestSuite):
    """export_pick_list_cost_excel"""

    def test_autofilter_scoped_to_header_and_item_rows(self):
        import openpyxl
        from io import BytesIO

        item_a = make_item()
        item_b = make_item()
        _seed_stock(item_a.name, qty=5, rate=10)
        _seed_stock(item_b.name, qty=5, rate=20)
        name = create_pick_list(COMPANY)
        add_item_to_pick_list(name, item_a.name, WAREHOUSE, 2)
        add_item_to_pick_list(name, item_b.name, WAREHOUSE, 3)

        export_pick_list_cost_excel(name)
        content = frappe.response["filecontent"]

        wb = openpyxl.load_workbook(BytesIO(content))
        ws = wb.active
        self.assertEqual(ws.auto_filter.ref, "A1:G3")

    def test_no_autofilter_when_pick_list_empty(self):
        import openpyxl
        from io import BytesIO

        name = create_pick_list(COMPANY)
        export_pick_list_cost_excel(name)
        content = frappe.response["filecontent"]

        wb = openpyxl.load_workbook(BytesIO(content))
        ws = wb.active
        self.assertEqual(ws.auto_filter.ref, None)

    def test_column_widths_survive_column_level_styles(self):
        import openpyxl
        from io import BytesIO

        item = make_item()
        _seed_stock(item.name, qty=5, rate=10)
        name = create_pick_list(COMPANY)
        add_item_to_pick_list(name, item.name, WAREHOUSE, 2)

        export_pick_list_cost_excel(name)
        content = frappe.response["filecontent"]

        wb = openpyxl.load_workbook(BytesIO(content))
        ws = wb.active
        # widths round-trip through xlsxwriter's internal unit conversion
        # with minor float drift, so assert comfortably above the
        # xlsxwriter-internal-default (~9.14) that this bug would produce,
        # not exact equality
        self.assertGreater(ws.column_dimensions["A"].width, 50)
        self.assertGreater(ws.column_dimensions["E"].width, 10)
        self.assertGreater(ws.column_dimensions["G"].width, 11)
        # alignment/format on those columns should also still be correct
        self.assertEqual(ws["A2"].alignment.horizontal, "left")


class TestBuildItemLink(ERPNextTestSuite):
    """_build_item_link"""

    def test_builds_url_with_display_text_and_desk_path(self):
        link = _build_item_link("SCX4IT1918", "Wire Locker 94-Inch")
        self.assertIsInstance(link, Url)
        self.assertEqual(link.text, "SCX4IT1918: Wire Locker 94-Inch")
        self.assertTrue(link._link.endswith("/desk/item/SCX4IT1918"))
        self.assertTrue(link._link.startswith("http"))

    def test_sanitizes_illegal_characters_in_display_text(self):
        link = _build_item_link("ITM001", "Bad\x00Name")
        self.assertEqual(link.text, "ITM001: BadName")

    def test_url_encodes_item_code_with_special_characters(self):
        link = _build_item_link("ITM 001/A", "Widget")
        self.assertIn("ITM%20001/A", link._link)


class TestAutofitColumnWidth(ERPNextTestSuite):
    """_autofit_column_width / _pixels_to_character_width"""

    def test_never_narrower_than_minimum(self):
        self.assertEqual(_autofit_column_width(["a"], 50), 50)

    def test_grows_beyond_minimum_for_long_text(self):
        result = _autofit_column_width(["A" * 100], 10)
        self.assertGreater(result, 10)

    def test_empty_values_returns_minimum(self):
        self.assertEqual(_autofit_column_width([], 20), 20)
